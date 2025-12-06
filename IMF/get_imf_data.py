import sys
import sdmx
from openpyxl import load_workbook
from datetime import datetime
import pandas as pd


def get_imf_data(workbook_path):
    """
    Lee parámetros del Excel y obtiene datos del IMF
    """
    
    print("=" * 60)
    print("Reading parameters from Excel...")
    print("=" * 60)
    
    try:
        # Load workbook and read Resume sheet
        wb = load_workbook(workbook_path, keep_vba=True, data_only=True)
        ws_resume = wb['Resume']
        
        # Read parameters from cells
        dataset = ws_resume['C23'].value
        country_name = ws_resume['C24'].value
        indicator_label = ws_resume['C26'].value
        start_date = str(ws_resume['E20'].value)
        end_date = str(ws_resume['E23'].value)
        
        print(f"Dataset: {dataset}")
        print(f"Country: {country_name}")
        print(f"Indicator: {indicator_label}")
        print(f"Period: {start_date} - {end_date}")
        
        # Get country code from Country_List sheet
        ws_country = wb['Country_List']
        country_code = None
        for row in ws_country.iter_rows(min_row=2, values_only=True):
            if row[1] == country_name:  # Column B = Country Name
                country_code = row[0]  # Column A = Country Code
                break
        
        print(f"Country Code: {country_code}")
        
        # Get indicator code from Indicator_List sheet
        ws_indicator = wb['Indicator_List']
        indicator_code = None
        for row in ws_indicator.iter_rows(min_row=2, values_only=True):
            if row[1] == indicator_label:  # Column B = Indicator Label
                indicator_code = row[0]  # Column A = Indicator Code
                break
        
        print(f"Indicator Code: {indicator_code}")
        
        if not all([dataset, country_code, indicator_code]):
            raise Exception("Missing required parameters")
        
        # Call IMF API
        print("\nCalling IMF API...")
        imf = sdmx.Client('IMF')
        
        key = f"{country_code}.{indicator_code}"
        print(f"Key: {key}")
        
        data_msg = imf.data(
            resource_id=dataset,
            key=key,
            params={'startPeriod': start_date, 'endPeriod': end_date}
        )
        
        # Convert to DataFrame
        df = sdmx.to_pandas(data_msg)
        
        if df is not None and not df.empty:
            print(f"Data received: {len(df)} observations")
            write_to_excel(wb, df, dataset, country_code, indicator_code)
            print("SUCCESS!")
        else:
            print("No data returned")
            write_error(wb, "No data found")
        
    except Exception as e:
        print(f"ERROR: {e}")
        import traceback
        traceback.print_exc()
        wb = load_workbook(workbook_path, keep_vba=True)
        write_error(wb, str(e))


def write_to_excel(wb, df, dataset, country_code, indicator_code):
    if 'API_Data' in wb.sheetnames:
        ws = wb['API_Data']
        ws.delete_rows(1, ws.max_row)
    else:
        ws = wb.create_sheet('API_Data')
    
    ws['A1'] = 'Dataset'
    ws['B1'] = dataset
    ws['A2'] = 'Country'
    ws['B2'] = country_code
    ws['A3'] = 'Indicator'
    ws['B3'] = indicator_code
    ws['A4'] = 'Retrieved'
    ws['B4'] = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    
    df_reset = df.reset_index()
    
    row = 6
    for col_idx, col_name in enumerate(df_reset.columns, start=1):
        ws.cell(row=row, column=col_idx, value=str(col_name))
    
    row = 7
    for idx, data_row in df_reset.iterrows():
        for col_idx, value in enumerate(data_row, start=1):
            ws.cell(row=row, column=col_idx, value=value)
        row += 1
    
    wb.save(wb.path)


def write_error(wb, error_msg):
    if 'API_Data' in wb.sheetnames:
        ws = wb['API_Data']
        ws.delete_rows(1, ws.max_row)
    else:
        ws = wb.create_sheet('API_Data')
    
    ws['A1'] = 'ERROR'
    ws['A2'] = error_msg
    wb.save(wb.path)


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("Usage: python script.py <workbook_path>")
        sys.exit(1)
    
    get_imf_data(sys.argv[1])


